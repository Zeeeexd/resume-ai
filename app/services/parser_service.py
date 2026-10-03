import re
from typing import Tuple
import fitz  # PyMuPDF
from docx import Document
from io import BytesIO
from fastapi import HTTPException, status


def extract_text_from_pdf(content: bytes) -> str:
    """Extract text from PDF bytes using PyMuPDF."""
    try:
        doc = fitz.open(stream=content, filetype="pdf")
        text_parts = []
        for page in doc:
            text_parts.append(page.get_text())
        doc.close()
        raw = "\n".join(text_parts)
        return _clean_text(raw)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Could not parse PDF: {str(e)}",
        )


def extract_text_from_docx(content: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    try:
        doc = Document(BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # also grab tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        raw = "\n".join(paragraphs)
        return _clean_text(raw)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Could not parse DOCX: {str(e)}",
        )


def _clean_text(text: str) -> str:
    """Remove excessive whitespace while preserving meaningful line breaks."""
    # collapse 3+ consecutive newlines to 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    # strip trailing spaces on each line
    lines = [line.rstrip() for line in text.split("\n")]
    text = "\n".join(lines)
    return text.strip()


def parse_resume(filename: str, content: bytes) -> Tuple[str, str]:
    """
    Returns (extracted_text, file_extension).
    Raises HTTPException for unsupported types or empty resumes.
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        text = extract_text_from_pdf(content)
    elif ext in ("docx", "doc"):
        text = extract_text_from_docx(content)
    else:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Only PDF and DOCX files are supported.",
        )

    if len(text.strip()) < 50:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="The uploaded document appears to be empty or unreadable.",
        )

    return text, ext
